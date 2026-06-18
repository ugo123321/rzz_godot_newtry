# -*- coding: utf-8 -*-
"""
[DEPRECATED 2026-06-08] 已被 tools/build_rewards_v6.py 取代。

旧脚本输出 config/json/upgrades_v6.json（仅 11 列原样转 JSON）。
新脚本输出 config/excel/rewards_v6.xlsx + config/json/rewards_v6.json（22 列规范化、抽取 weapon_mult / 标记 special_rule / extra_params 等）。

保留本文件作历史参考；运行它不会破坏新流水线，但产物 upgrades_v6.json 已无人使用。

将 ys构思_v6.xlsx 的 99 项 rewards_v6 转成 config/json/upgrades_v6.json。
非破坏性：不动现有的 config/json/upgrades.json。

v6 xlsx 列：name_cn, group, desc_cn, max level, notes, colour, id, dmg_layer, apply_type, trigger, pool_weight

输出 JSON 字段：
  id, name_cn, group, desc_cn, max_level, notes, colour, dmg_layer, apply_type, trigger, pool_weight,
  category, element  (派生)
"""
from __future__ import annotations
import io
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_XLSX = os.path.join(ROOT, "ys构思_v6.xlsx")
DST_JSON = os.path.join(ROOT, "config", "json", "upgrades_v6.json")

# group(中文) → category(英文标签，对齐 DamageInfo.category)
GROUP_TO_CATEGORY = {
    "基础属性": "physical",
    "生存防御": "survive",
    "普攻子弹": "bullet",
    "连击":     "combo",
    "画线轨迹": "trail",
    "强化球":   "orb",
    "环绕剑":   "sword",
    "召唤":     "summon",
    "元素":     "element",
}

# dmg_layer → element 派生(v6 用 E_FIRE/E_THUNDER/E_POISON/E_ICE 标四元素)
LAYER_TO_ELEMENT = {
    "E_FIRE":    "fire",
    "E_THUNDER": "thunder",
    "E_POISON": "poison",
    "E_ICE":    "ice",
}


def _strip(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v


def _category(group: str, dmg_layer: str) -> str:
    cat = GROUP_TO_CATEGORY.get(group, "")
    # 元素组按 dmg_layer 派生 element-* category，方便代码判分支
    if cat == "element" and dmg_layer in LAYER_TO_ELEMENT:
        cat = f"element_{LAYER_TO_ELEMENT[dmg_layer]}"
    return cat


def _element(dmg_layer: str) -> str:
    return LAYER_TO_ELEMENT.get(dmg_layer, "")


def main() -> int:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl 未安装。运行: pip install openpyxl", file=sys.stderr)
        return 2

    if not os.path.exists(SRC_XLSX):
        print(f"ERROR: 找不到源文件 {SRC_XLSX}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    if "rewards_v6" not in wb.sheetnames:
        print(f"ERROR: 工作簿没有 rewards_v6 表", file=sys.stderr)
        return 2

    ws = wb["rewards_v6"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("ERROR: rewards_v6 空表", file=sys.stderr)
        return 2

    headers = [str(h).strip() if h else "" for h in rows[0]]
    expected = ["name_cn", "group", "desc_cn", "max level", "notes", "colour",
                "id", "dmg_layer", "apply_type", "trigger", "pool_weight"]
    if headers != expected:
        print(f"ERROR: 列头不匹配。\n  期望: {expected}\n  实际: {headers}", file=sys.stderr)
        return 2

    out = []
    seen_ids = set()
    for i, r in enumerate(rows[1:], start=2):
        # 跳过完全空行
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        rec = dict(zip(headers, r))
        eid = _strip(rec["id"])
        if not eid:
            print(f"WARN: row {i} 没有 id，跳过", file=sys.stderr)
            continue
        if eid in seen_ids:
            print(f"WARN: row {i} 重复 id={eid}，跳过", file=sys.stderr)
            continue
        seen_ids.add(eid)
        group = _strip(rec["group"])
        dmg_layer = _strip(rec["dmg_layer"])
        out.append({
            "id":          eid,
            "name_cn":     _strip(rec["name_cn"]),
            "group":       group,
            "desc_cn":     _strip(rec["desc_cn"]),
            "max_level":   int(rec["max level"]) if rec["max level"] is not None else 1,
            "notes":       _strip(rec["notes"]),
            "colour":      _strip(rec["colour"]),
            "dmg_layer":   dmg_layer,
            "apply_type":  _strip(rec["apply_type"]),
            "trigger":     _strip(rec["trigger"]),
            "pool_weight": int(rec["pool_weight"]) if rec["pool_weight"] is not None else 0,
            "category":    _category(group, dmg_layer),
            "element":     _element(dmg_layer),
        })

    os.makedirs(os.path.dirname(DST_JSON), exist_ok=True)
    with open(DST_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    # 统计
    by_group: dict[str, int] = {}
    by_layer: dict[str, int] = {}
    by_element: dict[str, int] = {}
    for e in out:
        by_group[e["group"]] = by_group.get(e["group"], 0) + 1
        by_layer[e["dmg_layer"]] = by_layer.get(e["dmg_layer"], 0) + 1
        elk = e["element"] or "(none)"
        by_element[elk] = by_element.get(elk, 0) + 1

    print(f"✓ 写入 {DST_JSON}  ({len(out)} 项)")
    print("  by group:")
    for k, v in sorted(by_group.items(), key=lambda kv: -kv[1]):
        print(f"    {k}: {v}")
    print("  by dmg_layer:")
    for k, v in sorted(by_layer.items(), key=lambda kv: -kv[1]):
        print(f"    {k}: {v}")
    print("  by element:")
    for k, v in sorted(by_element.items(), key=lambda kv: -kv[1]):
        print(f"    {k}: {v}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
