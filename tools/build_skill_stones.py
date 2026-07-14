#!/usr/bin/env python3
"""Idempotently build/seed config/excel/skill_stones.xlsx (3 sheets).

Sheets:
  stones      : skill_id | enabled | decompose_gold | notes
  rules       : key | value | desc
  affix_stats : stat_key | name_cn | name_en

Seeding rule (matches CLAUDE.md §11 — never clobber hand-edited values):
  - `stones` rows are appended for every reward in config/json/rewards_v6.json
    whose pool_weight >= 1 AND whose skill_id is not already present.
    Existing rows (enabled / decompose_gold / notes the user edited) are
    NEVER overwritten or deleted.
  - `rules` rows are added only for keys that don't yet exist (existing values kept).
  - `affix_stats` rows are added only for stat_keys not yet present.

Run after rewards_v6.json changes to pick up newly poolable rewards:
    python tools/build_skill_stones.py
Then export:
    python tools/export_skill_stones_json.py
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "config" / "excel" / "skill_stones.xlsx"
REWARDS_JSON = ROOT / "config" / "json" / "rewards_v6.json"

HEADER_FILL = PatternFill("solid", fgColor="3a3a44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center")

# (key, value, desc) — defaults; only inserted if key is missing
DEFAULT_RULES = [
    ("small_drop_rate", 0.01, "小怪掉落技能石概率"),
    ("boss_drop_rate", 1.0, "boss掉落技能石概率(100%=1.0)"),
    ("rarity_white", 0.50, "掉落品质=白 概率"),
    ("rarity_blue", 0.30, "掉落品质=蓝 概率"),
    ("rarity_purple", 0.15, "掉落品质=紫 概率"),
    ("rarity_orange", 0.05, "掉落品质=橙 概率"),
    ("affix_count_white", 0, "白品质词条数"),
    ("affix_count_blue", 1, "蓝品质词条数"),
    ("affix_count_purple", 2, "紫品质词条数"),
    ("affix_count_orange", 3, "橙品质词条数"),
    ("affix_min", -0.10, "单词条属性浮动下限(-10%)"),
    ("affix_max", 0.10, "单词条属性浮动上限(+10%)"),
    ("decompose_gold_default", 5, "分解单块技能石默认金币(可被 stones.decompose_gold 覆盖)"),
    ("equipped_slot_count", 3, "技能石装备槽数量"),
]

# (stat_key, name_cn, name_en) — only inserted if stat_key missing
DEFAULT_AFFIX_STATS = [
    ("atk_pct", "攻击力", "ATK"),
    ("max_hp_pct", "生命", "HP"),
    ("ki_max_pct", "气力上限", "Max Stamina"),
    ("ki_regen_pct", "气力回复速度", "Stamina Regen"),
    ("crit_rate", "暴击率", "Crit Rate"),
    ("crit_damage", "暴击伤害", "Crit Damage"),
    ("move_speed_pct", "移速", "Move Speed"),
]


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
    ws.freeze_panes = "A2"


def _ensure_sheet(wb, name, headers):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, len(headers))
    return ws


def _existing_keys(ws, key_col=1):
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        k = row[0]
        if k is None or str(k).strip() == "":
            continue
        keys.add(str(k).strip())
    return keys


def _existing_skill_ids(ws):
    ids = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        k = row[0]
        if k is None or str(k).strip() == "":
            continue
        ids.add(str(k).strip())
    return ids


def build():
    # Load rewards to know which skill_ids are poolable (pool_weight >= 1)
    with open(REWARDS_JSON, "r", encoding="utf-8") as f:
        rewards = json.load(f)
    poolable_ids = []
    for r in rewards:
        try:
            pw = int(r.get("pool_weight", 0) or 0)
        except (TypeError, ValueError):
            pw = 0
        if pw >= 1:
            poolable_ids.append(str(r.get("id", "")))
    poolable_ids = [i for i in poolable_ids if i != ""]

    if XLSX.exists():
        wb = openpyxl.load_workbook(XLSX)
    else:
        wb = openpyxl.Workbook()
        # remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    # --- stones ---
    ws_stones = _ensure_sheet(wb, "stones", ["skill_id", "enabled", "decompose_gold", "notes"])
    have_stones = _existing_skill_ids(ws_stones)
    added_stones = 0
    next_row = ws_stones.max_row + 1
    for sid in poolable_ids:
        if sid in have_stones:
            continue
        ws_stones.cell(row=next_row, column=1, value=sid)
        ws_stones.cell(row=next_row, column=2, value=1)
        ws_stones.cell(row=next_row, column=3, value=5)
        ws_stones.cell(row=next_row, column=4, value="")
        next_row += 1
        added_stones += 1
    ws_stones.column_dimensions["A"].width = 28
    ws_stones.column_dimensions["B"].width = 9
    ws_stones.column_dimensions["C"].width = 16
    ws_stones.column_dimensions["D"].width = 30

    # --- rules ---
    ws_rules = _ensure_sheet(wb, "rules", ["key", "value", "desc"])
    have_rules = _existing_keys(ws_rules)
    added_rules = 0
    next_row = ws_rules.max_row + 1
    for key, value, desc in DEFAULT_RULES:
        if key in have_rules:
            continue
        ws_rules.cell(row=next_row, column=1, value=key)
        ws_rules.cell(row=next_row, column=2, value=value)
        ws_rules.cell(row=next_row, column=3, value=desc)
        next_row += 1
        added_rules += 1
    ws_rules.column_dimensions["A"].width = 24
    ws_rules.column_dimensions["B"].width = 12
    ws_rules.column_dimensions["C"].width = 40

    # --- affix_stats ---
    ws_affix = _ensure_sheet(wb, "affix_stats", ["stat_key", "name_cn", "name_en"])
    have_affix = _existing_keys(ws_affix)
    added_affix = 0
    next_row = ws_affix.max_row + 1
    for stat_key, cn, en in DEFAULT_AFFIX_STATS:
        if stat_key in have_affix:
            continue
        ws_affix.cell(row=next_row, column=1, value=stat_key)
        ws_affix.cell(row=next_row, column=2, value=cn)
        ws_affix.cell(row=next_row, column=3, value=en)
        next_row += 1
        added_affix += 1
    ws_affix.column_dimensions["A"].width = 18
    ws_affix.column_dimensions["B"].width = 16
    ws_affix.column_dimensions["C"].width = 18

    # order sheets: stones, rules, affix_stats
    order = ["stones", "rules", "affix_stats"]
    for i, name in enumerate(order):
        if name in wb.sheetnames:
            ws = wb[name]
            wb.move_sheet(ws, offset=i - wb.sheetnames.index(name))

    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    print(f"Saved {XLSX}")
    print(f"  stones: added {added_stones} new (poolable rewards = {len(poolable_ids)})")
    print(f"  rules:  added {added_rules} new")
    print(f"  affix_stats: added {added_affix} new")


if __name__ == "__main__":
    build()
