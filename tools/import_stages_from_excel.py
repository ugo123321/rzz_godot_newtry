#!/usr/bin/env python3
"""stages.xlsx -> stages.json 单向导入（安全）。

策划在 config/excel/stages.xlsx 的 `stages` sheet 里改完关卡配置后，跑这个脚本
把改动写回 config/json/stages.json，重启 Godot 即生效。

为什么不用 tools/export_config.py？
  export_config.py 是全量 excel->json，会顺带用硬编码的 2-boss dict 覆盖 bosses.json
  （现有 5 个 boss 会被洗成 2 个），还会因 monsters.xlsx 缺 elem_resist/vuln/name_en
  等列而洗掉 monsters.json 的 json-only 字段。本脚本只动 stages.json，零副作用。

跑法：python tools/import_stages_from_excel.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from export_config import (
    JSON_DIR,
    EXCEL_DIR,
    STAGE_HEADERS,
    normalize_stage_row,
    normalize_cell,
    save_json,
)

STAGES_XLSX = EXCEL_DIR / "stages.xlsx"
STAGES_SHEET = "stages"  # 主数据 sheet；另一个是「编码表」参考 sheet，不读


def _read_stages_sheet() -> list[dict]:
    if not STAGES_XLSX.exists():
        raise FileNotFoundError(f"找不到 {STAGES_XLSX}")
    wb = openpyxl.load_workbook(STAGES_XLSX, data_only=True)
    if STAGES_SHEET not in wb.sheetnames:
        raise RuntimeError(
            f"{STAGES_XLSX} 里没有 '{STAGES_SHEET}' sheet，"
            f"现有 sheets: {wb.sheetnames}"
        )
    ws = wb[STAGES_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out: list[dict] = []
    for raw in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        item: dict = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            val = normalize_cell(raw[i] if i < len(raw) else None)
            item[h] = "" if val is None else val
        if any(str(v).strip() != "" for v in item.values() if v is not None):
            out.append(item)
    return out


def main() -> None:
    rows = _read_stages_sheet()
    # normalize：reward_rooms 拆 list、空 room_type/theme 擦除、hp_coeff 转 float
    data = [normalize_stage_row(item) for item in rows]
    save_json("stages", data)
    print(f"[import_stages] {len(data)} stages <- {STAGES_XLSX}['{STAGES_SHEET}']")
    print("重启 Godot (F5) 生效。")


if __name__ == "__main__":
    main()
