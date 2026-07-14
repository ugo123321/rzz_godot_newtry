#!/usr/bin/env python3
"""Export config/excel/skill_stones.xlsx (3 sheets) + rewards_v6.json
-> config/json/skill_stones.json (runtime config for GameConfig).

Output schema:
  {
    "rules": { "<key>": <num/str>, ..., "affix_stats": ["atk_pct", ...] },
    "affix_display": { "atk_pct": {"name_cn":"攻击力","name_en":"ATK"}, ... },
    "stones": [
      { "skill_id":"...", "name_cn":"..", "name_en":"..", "rarity":"orange",
        "icon":"skill_46", "group":"普攻子弹", "desc_cn_game":"..", "desc_cn_game_en":"..",
        "decompose_gold": 5 }
    ]
  }

Only stones with enabled == 1 AND whose skill_id exists in rewards_v6.json are emitted.
name/rarity/icon/group/desc are joined from rewards_v6.json (single source of truth —
the xlsx only controls eligibility, decompose gold, and notes).

Run:
    python tools/export_skill_stones_json.py
Do NOT run build_rewards_v6_compact.py (would clobber rewards_v6_compact.xlsx hand-edits).
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC_XLSX = ROOT / "config" / "excel" / "skill_stones.xlsx"
REWARDS_JSON = ROOT / "config" / "json" / "rewards_v6.json"
OUT_JSON = ROOT / "config" / "json" / "skill_stones.json"


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "":
        return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return s


def _str(v):
    if v is None:
        return ""
    return str(v).strip()


def _to_bool_int(v, default=1):
    if v is None or v == "":
        return default
    if isinstance(v, bool):
        return 1 if v else 0
    try:
        return 1 if int(v) != 0 else 0
    except (TypeError, ValueError):
        s = str(v).strip().lower()
        return 1 if s in ("1", "true", "yes", "y", "on") else 0


def export():
    # rewards lookup by id (single source of truth for name/rarity/icon/...)
    with open(REWARDS_JSON, "r", encoding="utf-8") as f:
        rewards = json.load(f)
    by_id = {}
    for r in rewards:
        rid = str(r.get("id", ""))
        if rid != "":
            by_id[rid] = r

    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)

    # --- rules sheet -> dict ---
    rules = {}
    if "rules" in wb.sheetnames:
        ws = wb["rules"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            key = _str(row[0])
            if key == "":
                continue
            val = _num(row[1]) if len(row) > 1 else None
            rules[key] = val

    # defaults if missing
    rules.setdefault("small_drop_rate", 0.01)
    rules.setdefault("boss_drop_rate", 1.0)
    rules.setdefault("rarity_white", 0.50)
    rules.setdefault("rarity_blue", 0.30)
    rules.setdefault("rarity_purple", 0.15)
    rules.setdefault("rarity_orange", 0.05)
    rules.setdefault("affix_count_white", 0)
    rules.setdefault("affix_count_blue", 1)
    rules.setdefault("affix_count_purple", 2)
    rules.setdefault("affix_count_orange", 3)
    rules.setdefault("affix_min", -0.10)
    rules.setdefault("affix_max", 0.10)
    rules.setdefault("decompose_gold_default", 5)
    rules.setdefault("equipped_slot_count", 3)

    # --- affix_stats sheet -> display + ordered list ---
    affix_display = {}
    affix_stats_order = []
    if "affix_stats" in wb.sheetnames:
        ws = wb["affix_stats"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            stat_key = _str(row[0])
            if stat_key == "":
                continue
            cn = _str(row[1]) if len(row) > 1 else ""
            en = _str(row[2]) if len(row) > 2 else ""
            affix_display[stat_key] = {"name_cn": cn, "name_en": en}
            affix_stats_order.append(stat_key)
    if not affix_stats_order:
        # hard fallback so the game never runs with an empty affix pool
        for stat_key, cn, en in [
            ("atk_pct", "攻击力", "ATK"),
            ("max_hp_pct", "生命", "HP"),
            ("ki_max_pct", "气力上限", "Max Stamina"),
            ("ki_regen_pct", "气力回复速度", "Stamina Regen"),
            ("crit_rate", "暴击率", "Crit Rate"),
            ("crit_damage", "暴击伤害", "Crit Damage"),
            ("move_speed_pct", "移速", "Move Speed"),
        ]:
            affix_display[stat_key] = {"name_cn": cn, "name_en": en}
            affix_stats_order.append(stat_key)
    rules["affix_stats"] = affix_stats_order

    # --- stones sheet -> enriched list ---
    stones_out = []
    skipped = 0
    if "stones" in wb.sheetnames:
        ws = wb["stones"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            skill_id = _str(row[0])
            if skill_id == "":
                continue
            enabled = _to_bool_int(row[1], default=1) if len(row) > 1 else 1
            if enabled == 0:
                continue
            reward = by_id.get(skill_id)
            if reward is None:
                skipped += 1
                continue
            decompose_raw = _num(row[2]) if len(row) > 2 else None
            try:
                decompose_gold = int(decompose_raw) if decompose_raw is not None else int(rules.get("decompose_gold_default", 5))
            except (TypeError, ValueError):
                decompose_gold = int(rules.get("decompose_gold_default", 5))
            stones_out.append({
                "skill_id": skill_id,
                "name_cn": str(reward.get("name_cn", "")),
                "name_en": str(reward.get("name_en", "")),
                "rarity": str(reward.get("rarity", "")),
                "icon": str(reward.get("icon", "")),
                "group": str(reward.get("group", "")),
                "desc_cn_game": str(reward.get("desc_cn_game", "")),
                "desc_cn_game_en": str(reward.get("desc_cn_game_en", "")),
                "decompose_gold": decompose_gold,
            })

    out = {
        "rules": rules,
        "affix_display": affix_display,
        "stones": stones_out,
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"Wrote {OUT_JSON}")
    print(f"  stones: {len(stones_out)} enabled (skipped {skipped} unknown skill_ids)")
    print(f"  affix_stats: {len(affix_stats_order)}")
    return out


if __name__ == "__main__":
    export()
