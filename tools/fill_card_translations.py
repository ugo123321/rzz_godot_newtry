"""One-shot AI translation fill for remaining 90 cards. Run from project root."""
import openpyxl
from pathlib import Path

XLSX = Path('config/excel/rewards_v6_compact.xlsx')
wb = openpyxl.load_workbook(XLSX)
ws = wb['rewards']

T = {
    # Survival / Defense
    'sv_power_soul': ("Power Soul", "Max HP +15%/lv", "Max HP +15%"),
    'sv_desperate_heart': ("Desperate Heart", "Below 50% HP: Move +10%/lv", "Below 50% HP: Move +10%"),
    'sv_revive': ("Revive", "Revive once with 50% HP after death (once per stage)", "Revive once with 50% HP"),
    'sv_desperate_regen': ("Desperate Regen", "Below 30% HP: instant heal 30% +1%/lv (once per stage)", "Below 30% HP: instant heal 30%"),
    'sv_kill_revive': ("Battle Recovery", "On kill: 20% chance to heal 6%/lv max HP", "Kill chance to heal HP"),
    'sv_stand_guard': ("Mountain Stance", "Standing still 1.5s: Defense +30%/lv (max 50%)", "Standing still grants defense"),
    'sv_angel_shelter': ("Angel's Shelter", "Max HP +10%/lv. Invincible 1.5s after being hit (CD 6s)", "Max HP +10%, invincible briefly after hit"),
    'sv_blood_power': ("Blood Power", "Max HP +10%/lv", "Max HP +10%"),
    'sv_demon_recover': ("Demon Recovery", "On kill: 15% chance to heal 5%/lv max HP", "Kill chance to heal HP"),
    'sv_life_spring': ("Life Spring", "Max HP +8%/lv. Heal 30% on stage clear", "Max HP +8%, heal on stage clear"),

    # Bullet
    'bullet_storm_king': ("Storm King", "Bullets +3, Attack +15%", "Bullets +3, Attack +15%"),
    'bullet_spirit_bomb': ("Spirit Bomb", "Bullets become spirit bombs: farther flight = more damage (+50% max)", "Bullets travel farther for higher damage"),
    'bullet_swift_shoot': ("Swift Shoot", "Atk Speed +30%/lv, Attack -10%/lv", "Atk Speed +30%, Attack -10%"),
    'bullet_homing': ("Homing Soul", "Bullets home in. Attack -5%/lv", "Bullets home in. Attack -5%"),
    'bullet_fire_support': ("Fire Support", "Attack +10%/lv. Hits sometimes trigger explosion (AOE 80% ATK)", "Attack +10%, hits sometimes explode"),
    'bullet_split': ("Splitter", "On hit: split into 3 small bullets (50% ATK, +10%/lv)", "On hit: split into 3 small bullets"),
    'bullet_bounce': ("Bouncing Bullet", "Bounce +1/lv (each bounce 0.6x damage)", "Bullets gain bounce"),
    'bullet_mirror': ("Mirror Bullet", "After hit, bullet returns dealing 0.6 ATK along the way", "Bullets return after hitting"),
    'bullet_side': ("Side Shot", "Fire 2 extra bullets at +/-30 degrees", "Fire extra angled bullets"),
    'bullet_count': ("Multi Bullet", "Bullets per shot +1/lv", "Bullets +1"),
    'bullet_beam': ("Light Beam", "Attack +10%/lv. Sometimes fires piercing beam (100% ATK)", "Attack +10%, sometimes piercing beam"),

    # Combo
    'combo_multi': ("Multi Slash", "Each slash creates 2 ninja clones", "Each slash creates 2 ninja clones"),
    'combo_black_hole': ("Black Hole", "Every 8 combo: spawn black hole pulling enemies (2.0 ATK). Combo -2", "Every 8 combo: spawn black hole"),
    'combo_fireball': ("Greater Fireball", "Every 14 combo: launch fireball (2.5 ATK). Combo -2", "Every 14 combo: launch fireball"),
    'combo_water_tornado': ("Water Tornado", "Every 6 combo: launch tornado (2.2 ATK). Combo -1", "Every 6 combo: launch tornado"),
    'combo_blade_storm': ("Blade Storm", "Every 8 combo: 6 blades storm (2.0 ATK). Combo +1", "Every 8 combo: blade storm"),
    'combo_charge': ("Charge Slash", "Stand still 0.4s before slash: damage +20%/lv (max x1.5)", "Pre-slash standing boosts damage"),
    'combo_thunder': ("Thunder Strike", "Every 8 combo: lightning bolt (3.0 ATK). Combo -1", "Every 8 combo: lightning bolt"),
    'combo_shuriken': ("Spirit Shuriken", "At combo cap: slash end spawns +2 shuriken (0.6 ATK each)", "At combo cap: spawn shuriken"),
    'combo_crit': ("Combo Crit", "Per 10 combo: next slash Crit Dmg +10%/lv", "Per 10 combo: next crit damage +10%"),

    # Trail
    'trail_multi': ("Multi Trail", "When drawing: spawn 2 parallel trails", "When drawing: spawn 2 parallel trails"),
    'trail_fire_wall': ("Fire Wall", "Trail leaves flame wall 5s +1s/lv (0.8 ATK/0.5s). Bullets ignore wall", "Trail leaves flame wall; bullets pass through"),
    'trail_thunder_field': ("Thunder Field", "Trail spawns thunder field 3s +1s/lv (stun + 0.5 ATK every 0.5s)", "Trail spawns thunder field"),
    'trail_poison_fog': ("Corrosive Fog", "Trail spawns poison fog 4s +1s/lv (0.35 ATK/s, +25% dmg taken)", "Trail spawns poison fog"),
    'trail_frost': ("Frost Trail", "Trail spawns frost 3s +1s/lv (slow 40%, 0.5 ATK every 0.5s)", "Trail spawns frost"),
    'trail_slash_wave': ("Slash Wave", "At combo cap: slash end creates AOE wave (2.5 ATK, 200px +0.3/lv)", "At combo cap: slash creates wave"),
    'trail_loop_explode': ("Loop Explode", "At combo cap: first closed-loop trail explodes (3.0 ATK +0.3/lv)", "Closed-loop trail explodes at combo cap"),
    'trail_pierce': ("Trail Pierce", "Trails pierce obstacles", "Trails pierce obstacles"),
    'trail_width': ("Trail Width", "Trail width +10%/lv", "Trail width +10%"),
    'trail_dmg': ("Trail Power", "Slash damage +8%/lv", "Slash damage +8%"),

    # Orb
    'orb_tide': ("Tidal Orb", "Spawn 1 power orb every 8s (-2s/lv)", "Spawn power orb every 8s (-2s/lv)"),
    'orb_mark': ("Orb Mark", "On pickup: 30% chance to mark target. Next hit +5%/lv", "Pickup may mark target for bonus damage"),
    'orb_field': ("Field Boost", "Per orb on field: damage +30%/lv", "Per orb on field: damage +30%"),
    'orb_magnet': ("Magnet", "Pickup range +30%/lv. Crit dmg +30%/lv inside aura", "Pickup range +30%, crit dmg +30% in aura"),
    'orb_burst': ("Burst Pickup", "On pickup: explode at orb position (1.5 ATK +0.3/lv)", "On pickup: explode at position (1.5 ATK)"),
    'orb_glow': ("Subtle Glow", "On pickup: bullet/slash dmg +15%/lv for 6s", "On pickup: damage +15% for 6s"),
    'orb_fire': ("Fire Orb", "Element orb. Pickup AOE burn 1.0 ATK (+0.2/lv)", "Element orb. Pickup AOE burn"),
    'orb_ice': ("Ice Orb", "Element orb. Pickup AOE freeze 2.0s (+0.2s/lv)", "Element orb. Pickup AOE freeze"),
    'orb_poison': ("Poison Orb", "Element orb. Pickup AOE poison 0.6 ATK/s for 3s", "Element orb. Pickup AOE poison"),
    'orb_thunder': ("Thunder Orb", "Element orb. Pickup AOE thunder 1.5 ATK, chain +1/lv", "Element orb. Pickup AOE thunder chain"),

    # Sword
    'sword_double': ("Twin Swords", "Orbiting swords x2", "Orbiting swords x2"),
    'sword_rage': ("Raging Sword", "Each sword per cycle: 15% chance to launch slash wave (2.5 ATK, +5%/lv area)", "Sword sometimes launches slash wave"),
    'sword_guard': ("Guard Sword", "+2 guard swords that block ranged bullets", "2 guard swords block ranged bullets"),
    'sword_length': ("Long Reach", "Orbit radius +15%/lv", "Orbit radius +15%"),
    'sword_speed': ("Spin Boost", "Orbit speed +15%/lv", "Orbit speed +15%"),
    'sword_blood': ("Vampire Sword", "+2 vampire swords. 1% lifesteal on hit (CD 0.6s)", "2 vampire swords with lifesteal"),
    'sword_dmg': ("Razor Edge", "Sword damage +15%/lv", "Sword damage +15%"),
    'sword_flame': ("Flame Sword", "+1 flame sword (0.40 ATK burn/0.5s, 2s duration)", "Flame Sword"),
    'sword_thunder': ("Thunder Sword", "+1 thunder sword (0.40 ATK chain damage)", "Thunder Sword"),
    'sword_poison': ("Poison Sword", "+1 poison sword (0.30 ATK poison/s for 3s)", "Poison Sword"),
    'sword_frost': ("Frost Sword", "+1 frost sword (0.30 ATK damage, slow 30% 1s)", "Frost Sword"),

    # Summon
    'summon_pact': ("Summon Pact", "Gain Black Ninja. Summon damage +50%, Atk Spd +15%, Move +50%", "Black Ninja summoned. Summon dmg +50%, Atk Spd +15%, Move +50%"),
    'summon_rage': ("Summon Rage", "Summon dmg +12%/lv. Player dmg -12%/lv", "Summon dmg +12%, Player dmg -12%"),
    'summon_king': ("Demon King", "Summon demon king (ranged AOE attack)", "Summon demon king"),
    'summon_god': ("Thunder God", "Summon thunder god (ranged single-target)", "Summon thunder god"),
    'summon_gorilla': ("Gorilla", "Summon gorilla (ranged. Every 10s throws stunning banana, 3s stun)", "Summon gorilla"),
    'summon_dmg': ("Summon Power", "Summon damage +10%/lv", "Summon damage +10%"),
    'summon_thunder': ("Thunder Summon", "Summon thunder. Every 3s strikes random nearby AOE", "Summon thunder"),
    'summon_bear': ("Bear Summon", "Summon bear ally (ranged single-target)", "Summon bear"),
    'summon_snake': ("Snake Summon", "Summon snake ally (ranged single-target)", "Summon snake"),
    'summon_fire': ("Fire Spirit", "Summon fire spirit ally (ranged single-target)", "Summon fire spirit"),

    # Element
    'elem_fire_plus': ("Fire+", "Attack +30%/lv. Burn frequency +15%/lv", "Attack +30%, Burn rate +15%"),
    'elem_thunder_plus': ("Thunder+", "Attack +30%/lv. Chain targets +1/lv", "Attack +30%, Chain targets +1"),
    'elem_poison_plus': ("Poison+", "Attack +30%/lv. 8% chance to spread poison +8%/lv", "Attack +30%, 8% poison spread +8%"),
    'elem_ice_plus': ("Ice+", "Attack +30%/lv. 8% chance to spread freeze +8%/lv", "Attack +30%, 8% freeze spread +8%"),
    'elem_fire_bullet': ("Fire Bullet", "Bullets become fire: hits apply burn", "Bullets become fire bullets"),
    'elem_thunder_bullet': ("Thunder Bullet", "Bullets become thunder: hits apply chain lightning", "Bullets become thunder bullets"),
    'elem_poison_bullet': ("Poison Bullet", "Bullets become poison: hits apply poison", "Bullets become poison bullets"),
    'elem_ice_bullet': ("Ice Bullet", "Bullets become ice: hits apply slow + freeze", "Bullets become ice bullets"),

    # Demon
    'demon_scythe': ("Death Scythe", "At combo cap: throw piercing scythe from slash end (4 ATK). Max HP -30%", "At combo cap: throw piercing scythe. Max HP -30%"),
    'demon_sulfur_laser': ("Sulfur Laser", "Every 6s: fire 1s piercing laser that burns. Max HP -30%", "Periodic sulfur laser. Max HP -30%"),
    'demon_baby': ("Demon Baby", "Summon Demon Baby (ranged piercing laser). Max HP -30%", "Summon Demon Baby. Max HP -30%"),
    'demon_nine_lives': ("Nine Lives Cat", "Start with 1 life +8 extra revives (HP=1). Max HP -99%", "Multiple revives. Max HP -99%"),
    'demon_vampire': ("Vampire", "On kill: 5% chance to heal 20% max HP. Max HP -10%", "Kill chance to heal. Max HP -10%"),
    'demon_blood_blade': ("Blood Knife", "Replace basic shot with Blood Knife (pierce + multishot x2 + range +100%). Atk Spd -80%, Move -20%", "Replace basic with Blood Knife"),

    # Angel
    'angel_holy_bullet': ("Holy Bullet", "Basic attacks have 10% chance to summon holy AOE strike", "Basic attacks summon holy strike"),
    'angel_light_ward': ("Light Ward", "Damage taken -50%", "Damage taken -50%"),
    'angel_baby': ("Angel Baby", "Summon Angel Baby (ranged single-target ally)", "Summon Angel Baby"),
    'angel_fate_spear': ("Spear of Fate", "Spear of Fate x1", "Spear of Fate"),
    'angel_proximity_slow': ("Limit Break Technique", "Enemies within 300px slowed by distance (max 50%)", "Nearby enemies slowed"),
}

id_to_row = {}
for row_idx in range(3, ws.max_row + 1):
    rid = ws.cell(row=row_idx, column=1).value
    if rid:
        id_to_row[str(rid)] = row_idx

missing = []
count = 0
for rid, (n, d, dg) in T.items():
    if rid not in id_to_row:
        missing.append(rid)
        continue
    r = id_to_row[rid]
    ws.cell(row=r, column=36, value=n)
    ws.cell(row=r, column=37, value=d)
    ws.cell(row=r, column=38, value=dg)
    count += 1

wb.save(XLSX)
print(f"Wrote {count} translations")
if missing:
    print("MISSING:", missing)
