#!/usr/bin/env python3
"""Push config/json into Excel. Run after editing JSON; before export_config for stages/monsters only.

覆盖：chapters / monsters / stages(+编码表) / upgrade_fx / rewards_v6 / game_tuning。
game_tuning 的 json 是运行时权威源（game_config.gd 读 game_tuning.json），xlsx 是它的
策划可读镜像——历史上有 20 个 key 只加在 json 没回写 xlsx 导致 xlsx stale，现在由本脚本
每次从 json 重建 xlsx，保持完整镜像。"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from export_config import (
    CHAPTER_HEADERS,
    EXCEL_DIR,
    JSON_DIR,
    MONSTER_HEADERS,
    REWARDS_V6_HEADERS,
    STAGE_HEADERS,
    TUNING_HEADERS,
    UPGRADE_FX_HEADERS,
    write_sheet,
)


def _load_json(name: str) -> list[dict]:
    with (JSON_DIR / f"{name}.json").open(encoding="utf-8") as f:
        return json.load(f)


def _monster_row(item: dict) -> list:
    return [item.get(h, "") for h in MONSTER_HEADERS]


def _stage_row(item: dict) -> list:
    row: list = []
    for h in STAGE_HEADERS:
        val = item.get(h, "")
        if h == "reward_rooms" and isinstance(val, list):
            val = ",".join(str(x) for x in val)
        row.append(val)
    return row


def _chapter_row(item: dict) -> list:
    return [item.get(h, "") for h in CHAPTER_HEADERS]


def _upgrade_fx_row(item: dict) -> list:
    return [item.get(h, "") for h in UPGRADE_FX_HEADERS]


def _tuning_row(item: dict) -> list:
    return [item.get(h, "") for h in TUNING_HEADERS]


def _reward_v6_row(item: dict) -> list:
    """rewards_v6.json 里 extra_params 是 dict，要回写成 JSON 字符串再进 Excel 单元格。"""
    row: list = []
    for h in REWARDS_V6_HEADERS:
        val = item.get(h, "")
        if h == "extra_params" and isinstance(val, dict):
            val = json.dumps(val, ensure_ascii=False, sort_keys=True) if val else ""
        row.append(val)
    return row


def _rewrite_sheet(path: Path, headers: list[str], rows: list[list], sheet_title: str | None = None) -> None:
    if path.exists():
        path.unlink()
    wb = Workbook()
    ws = wb.active
    if sheet_title:
        ws.title = sheet_title
    write_sheet(ws, headers, rows)
    wb.save(path)
    print(f"Excel: {path}")


# stages.xlsx 专属：主 sheet + 「编码表」参考 sheet 同写一个 workbook，
# 这样每次 sync 都会重建编码表，且小怪/boss 中文名随 json 同步。
# stages.xlsx 的 26 列 schema 见 export_config.STAGE_HEADERS。
_ENCODING_HEADERS = ["字段", "可填值", "中文名", "说明"]

# stages 表小怪列名 -> monsters.json kind_id
_MONSTER_FIELD_TO_KIND = [
    ("normal", "NORMAL"), ("elite", "ELITE"), ("shield", "SHIELD"), ("berserker", "BERSERKER"),
    ("splitter", "SPLITTER"), ("archer", "ARCHER"), ("fire_mage", "FIRE_MAGE"),
    ("shotgun", "SHOTGUN"), ("cross_shooter", "CROSS_SHOOTER"), ("bounce_slime", "BOUNCE_SLIME"),
    ("snake_shooter", "SNAKE_SHOOTER"), ("jumper", "JUMPER"), ("laser", "LASER"),
    ("mini_centipede", "MINI_CENTIPEDE"), ("dasher", "DASHER"), ("teleporter", "TELEPORTER"),
]


def _encoding_rows(monsters: list[dict], bosses: dict) -> list[list]:
    monster_by_id = {str(m.get("kind_id", "")): m for m in monsters}
    rows: list[list] = []
    # room_type
    rows.append(["room_type", "(空)", "普通战斗", "默认值，普通战斗关"])
    rows.append(["room_type", "attr_forge", "属性打造关", "进入打造房（无需天赋卡解锁）"])
    rows.append(["room_type", "reward", "奖励房", "走 reward_rooms 抽房"])
    # reward_rooms
    rows.append(["reward_rooms", "wheel", "转盘", "当前唯一实现的奖励房类型"])
    # theme
    rows.append(["theme", "(空)", "普通关", "无主题"])
    rows.append(["theme", "demon", "恶魔关", "通关弹恶魔主题奖励卡（接受/放弃）"])
    rows.append(["theme", "angel", "天使关", "通关弹天使主题奖励卡（接受/放弃）"])
    rows.append(["theme", "random", "恶魔/天使随机", "本关随机分配恶魔或天使主题，本局内固定"])
    # boss_id
    rows.append(["boss_id", "(空)", "无 boss", "普通战斗关"])
    for bid, b in bosses.items():
        rows.append(["boss_id", str(bid), str(b.get("name", "")), ""])
    # 小怪列 -> kind_id/中文名
    for field, kind in _MONSTER_FIELD_TO_KIND:
        m = monster_by_id.get(kind, {})
        rows.append(["小怪列 " + field, kind, str(m.get("name_cn", "")), "填该关此种怪物的生成数量"])
    # hp_coeff
    rows.append(["hp_coeff", "1.0 / 1.05 / ...", "每关 HP 系数",
                 "累积乘积：第N关 HP = base × ∏(coeff[0..N])，取代旧 pow(stage_hp_growth) 指数曲线"])
    return rows


def _write_stages_workbook(stages: list[dict], monsters: list[dict], bosses: dict) -> None:
    path = EXCEL_DIR / "stages.xlsx"
    if path.exists():
        path.unlink()
    wb = Workbook()
    ws = wb.active
    ws.title = "stages"
    write_sheet(ws, STAGE_HEADERS, [_stage_row(s) for s in stages])
    enc = wb.create_sheet("编码表")
    write_sheet(enc, _ENCODING_HEADERS, _encoding_rows(monsters, bosses))
    wb.save(path)
    print(f"Excel: {path} (+编码表 sheet)")


def main() -> None:
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    chapters = _load_json("chapters")
    monsters = _load_json("monsters")
    stages = _load_json("stages")
    bosses = _load_json("bosses")  # dict
    _rewrite_sheet(
        EXCEL_DIR / "chapters.xlsx",
        CHAPTER_HEADERS,
        [_chapter_row(c) for c in chapters],
    )
    _rewrite_sheet(
        EXCEL_DIR / "monsters.xlsx",
        MONSTER_HEADERS,
        [_monster_row(m) for m in monsters],
    )
    _write_stages_workbook(stages, monsters, bosses)
    upgrade_fx = _load_json("upgrade_fx")
    _rewrite_sheet(
        EXCEL_DIR / "upgrade_fx.xlsx",
        UPGRADE_FX_HEADERS,
        [_upgrade_fx_row(u) for u in upgrade_fx],
    )
    # game_tuning: json 权威，xlsx 镜像（sheet 名保持 game_tuning）
    game_tuning = _load_json("game_tuning")
    _rewrite_sheet(
        EXCEL_DIR / "game_tuning.xlsx",
        TUNING_HEADERS,
        [_tuning_row(t) for t in game_tuning],
        sheet_title="game_tuning",
    )
    # rewards_v6: JSON 为权威源，回写到 Excel 便于策划编辑
    rewards_v6_path = JSON_DIR / "rewards_v6.json"
    if rewards_v6_path.exists():
        rewards_v6 = _load_json("rewards_v6")
        _rewrite_sheet(
            EXCEL_DIR / "rewards_v6.xlsx",
            REWARDS_V6_HEADERS,
            [_reward_v6_row(r) for r in rewards_v6],
        )
    print("Done.")


if __name__ == "__main__":
    main()
