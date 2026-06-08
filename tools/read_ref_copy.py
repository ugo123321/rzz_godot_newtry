import openpyxl
wb = openpyxl.load_workbook("ref_copy.xlsx")
print("Sheets:", wb.sheetnames)
for n in wb.sheetnames:
    ws = wb[n]
    print(f"\n=== {n} ===")
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 100), values_only=True):
        print(r)