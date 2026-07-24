#!/usr/bin/env python3
"""Convert config/excel/equipments.xlsx Sheet1 -> config/json/equipments.json

Schema (per equipment):
    {
      "def_id": "iron_dagger",
      "name_cn": "铁制短刀", "name_en": "Iron Dagger",
      "slot": "weapon",
      "icon_path": "res://assets/.../icon_equip_iron_dagger.png",
      "is_rare": 0,
      "per_level_bonuses": {"attack": 1},   # G 列：每次升级叠加的绝对值（每装备一份）
      "tiers": [
        {"quality": 0, "stat_bonuses": {...}, "flag": null},                 # 数值型无 desc
        {"quality": 1, "stat_bonuses": {...}, "flag": null},
        {"quality": 2, "stat_bonuses": {...}, "flag": null},
        {"quality": 3, "stat_bonuses": {}, "flag": "bullet_homing",
         "effect_desc_cn": "...", "effect_desc_en": "..."},                 # flag 保留 desc
      ]
    }

Columns A-G: 名称 / 是否稀有 / 部位 / icon / 品质 / 实际效果(F) / 升级每级实际效果(G).
数值型属性 stat_bonuses 用绝对值 key（attack/max_hp/crit_rate/crit_damage/move_speed/
max_ki/ki_regen/invincible_time/ki_per_pixel）。显示百分比由 lobby_state 运行时算。

Skip the "草稿，无视此sheet" sheet. Row-groups of 4 rows per equipment (name
repeats in column A). Runtime consumer: scripts/autoload/lobby_state.gd.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC_XLSX = ROOT / "config" / "excel" / "equipments.xlsx"
OUT_JSON = ROOT / "config" / "json" / "equipments.json"

# 名称（xlsx 中文名）→ def_id / 英文名
NAME_TO_META = {
    "铁制短刀":   ("iron_dagger",        "Iron Dagger"),
    "轻布甲":     ("light_cloth_armor",  "Light Cloth Armor"),
    "硬木鞋":     ("hard_wood_shoes",    "Hard Wood Shoes"),
    "毛绒帽":     ("fur_hat",            "Fur Hat"),
    "暴风大剑":   ("storm_greatsword",   "Storm Greatsword"),
    "轻灵之靴":   ("nimble_boots",       "Nimble Boots"),
    "丛林甲":     ("jungle_armor",       "Jungle Armor"),
    "坚固头盔":   ("sturdy_helmet",      "Sturdy Helmet"),
}

# 部位（xlsx）→ slot key（游戏内）
SLOT_CN_TO_KEY = {
    "武器": "weapon",
    "衣服": "armor",
    "鞋":   "shoes",
    "头盔": "helmet",
    "项链": "necklace",
    "戒指": "ring",
}

# 品质（xlsx 单字）→ 品质代码
QUALITY_CN_TO_CODE = {"白": 0, "蓝": 1, "紫": 2, "橙": 3}

# stat_key → 该 stat 是否「倍率型」（百分比加成，旧 pct 语义）。
# v6 改版后全部走绝对值；crit_damage 也改为绝对加（不再 pct 乘）。
# 这里 EFFECT_PATTERNS 只负责把中文文案解析成 (stat_bonuses, flag, effect_desc_en)。
# 数值型 tier 不再需要 effect_desc_en（运行时算百分比），只 flag 保留 desc。
EFFECT_PATTERNS = [
    # 攻击力 +N
    (re.compile(r"^攻击力\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"attack": float(m.group(1))}, None, "")),
    # 生命 +N（浮点心数，如 0.5）
    (re.compile(r"^生命\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"max_hp": float(m.group(1))}, None, "")),
    # 暴击率 +N（0–1 浮点，如 0.05）
    (re.compile(r"^暴击率\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"crit_rate": float(m.group(1))}, None, "")),
    # 暴击伤害倍率基础值 +N（绝对加成，如 0.3 / 1）
    (re.compile(r"^暴击伤害倍率基础值\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"crit_damage": float(m.group(1))}, None, "")),
    # 移动速度 +N
    (re.compile(r"^移动速度\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"move_speed": float(m.group(1))}, None, "")),
    # 气力上限 +N（绝对，新 key；旧是 max_ki_pct）
    (re.compile(r"^气力上限\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"max_ki": float(m.group(1))}, None, "")),
    # 气力回复速度 +N（绝对，新 key；旧是 ki_regen_pct）
    (re.compile(r"^气力回复速度\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"ki_regen": float(m.group(1))}, None, "")),
    # 受击无敌时间 +N（全新）
    (re.compile(r"^受击无敌时间\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"invincible_time": float(m.group(1))}, None, "")),
    # 划线气力消耗 +N（可为负；负=降消耗）
    (re.compile(r"^划线气力消耗\+(-?\d+(?:\.\d+)?)$"),
     lambda m: ({"ki_per_pixel": float(m.group(1))}, None, "")),
    # 4 个 flag（纯机制，保留 desc）
    (re.compile(r"^子弹获得追踪效果$"),
     lambda m: ({}, "bullet_homing", "Bullets seek nearby enemies")),
    (re.compile(r"^持续电击靠近的敌人$"),
     lambda m: ({}, "shock_aura", "Continuously shocks nearby enemies")),
    (re.compile(r"^每关随机刷出的树木数量翻倍$"),
     lambda m: ({}, "tree_x2", "Trees spawned per stage x2")),
    (re.compile(r"^受击时(\d+)%概率免伤$"),
     lambda m: ({}, "hit_dodge_5pct",
                f"{m.group(1)}% chance to negate incoming hits")),
]


def parse_effect(text: str) -> tuple[dict, str | None, str]:
    """(stat_bonuses, flag, effect_desc_en) —— 数值型 desc 留空（运行时算百分比）；
    flag 型保留英文 desc。找不到匹配就报错，避免静默失败。"""
    text = (text or "").strip()
    for pat, builder in EFFECT_PATTERNS:
        m = pat.match(text)
        if m:
            return builder(m)
    raise ValueError(f"Unknown effect text: {text!r}")


def icon_path_for(icon_key: str) -> str:
    return f"res://assets/ui/icons/equipment/{icon_key}.png"


def convert():
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    if "Sheet1" not in wb.sheetnames:
        raise RuntimeError(f"Sheet1 missing in {SRC_XLSX}")
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))[1:]  # skip header row

    equipments: dict[str, dict] = {}  # def_id -> record
    order: list[str] = []             # keep insertion order for stable output
    per_level_seen: dict[str, dict] = {}  # def_id -> per_level_bonuses（校验 4 行一致）

    for r in rows:
        if not r or r[0] is None:
            continue
        name_cn = str(r[0]).strip()
        if name_cn not in NAME_TO_META:
            raise ValueError(f"Unknown equipment name: {name_cn!r}")
        def_id, name_en = NAME_TO_META[name_cn]
        is_rare = int(r[1] or 0)
        slot_cn = str(r[2] or "").strip()
        slot_key = SLOT_CN_TO_KEY.get(slot_cn)
        if slot_key is None:
            raise ValueError(f"Unknown slot: {slot_cn!r}")
        icon_key = str(r[3] or "").strip()
        if not icon_key:
            raise ValueError(f"{def_id}: missing icon slug in column D")
        quality_cn = str(r[4] or "").strip()
        if quality_cn not in QUALITY_CN_TO_CODE:
            raise ValueError(f"Unknown quality: {quality_cn!r}")
        quality = QUALITY_CN_TO_CODE[quality_cn]
        effect_cn = str(r[5] or "").strip()
        stat_bonuses, flag, effect_en = parse_effect(effect_cn)
        # G 列：升级每级实际效果（每装备一份，4 行相同）
        per_level_cn = str(r[6] or "").strip()
        per_level_bonuses, _pl_flag, _pl_en = parse_effect(per_level_cn)
        if _pl_flag is not None:
            raise ValueError(f"{def_id}: per-level effect must be a stat, got flag: {per_level_cn!r}")

        if def_id not in equipments:
            equipments[def_id] = {
                "def_id": def_id,
                "name_cn": name_cn,
                "name_en": name_en,
                "slot": slot_key,
                "icon_path": icon_path_for(icon_key),
                "is_rare": is_rare,
                "tiers": [None, None, None, None],
                "per_level_bonuses": per_level_bonuses,
            }
            order.append(def_id)
            per_level_seen[def_id] = per_level_bonuses
        else:
            # 校验 4 行的 G 列一致
            if per_level_bonuses != per_level_seen[def_id]:
                raise ValueError(
                    f"{def_id}: per-level effect inconsistent across quality rows: "
                    f"{per_level_bonuses} vs {per_level_seen[def_id]}")
        rec = equipments[def_id]
        if rec["tiers"][quality] is not None:
            raise ValueError(
                f"Duplicate tier for {def_id} quality={quality}")
        tier_rec = {
            "quality": quality,
            "stat_bonuses": stat_bonuses,
            "flag": flag,
        }
        # 数值型 tier 不写 desc（运行时算百分比）；flag tier 保留 desc
        if flag is not None:
            tier_rec["effect_desc_cn"] = effect_cn
            tier_rec["effect_desc_en"] = effect_en
        rec["tiers"][quality] = tier_rec

    # Validate all 4 tiers present
    for def_id in order:
        rec = equipments[def_id]
        for q in range(4):
            if rec["tiers"][q] is None:
                raise ValueError(
                    f"{def_id} missing tier quality={q}")

    out = {"equipments": [equipments[d] for d in order]}
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"Wrote {OUT_JSON} ({len(out['equipments'])} equipments)")
    return out


if __name__ == "__main__":
    convert()
